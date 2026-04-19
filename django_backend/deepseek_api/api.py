import json
import logging
import os
import re
import zipfile
import time
import requests
from typing import Generator
from asgiref.sync import sync_to_async
from ninja import File, NinjaAPI, Router, Schema

from django.conf import settings
from django.db import connection
from django.http import HttpResponse, StreamingHttpResponse
from ninja import File, NinjaAPI, Router
from ninja.files import UploadedFile as NinjaUploadedFile

from . import services
from .agents import LlmConfig, MultiAgentConfig, Orchestrator, OrchestratorInput
from .dashboard_stats import build_dashboard_stats
from .models import APIKey
from .query_service import (
    export_query_records,
    get_query_facets,
    get_query_record_detail,
    list_query_records,
)
from .schemas import (
    ChatIn,
    ErrorResponse,
    HistoryOut,
    LoginIn,
    LoginOut,
    RuntimeConfigOut,
    SessionListOut,
    SessionRenameIn,
)
from .services import (
    delete_conversation_session,
    get_or_create_session,
    list_user_sessions,
    model_api_call,
    rename_conversation_session,
)

logger = logging.getLogger(__name__)

api = NinjaAPI(title="DeepSOC——基于多智能体协同与RAG架构的智能安全运营系统 API", version="0.0.1")


def api_key_auth(request):
    """验证请求头中的 API Key。"""
    auth_header = request.headers.get("Authorization")
    if not auth_header:
        return None
    try:
        scheme, key = auth_header.split()
        if scheme.lower() != "bearer":
            return None
        api_key = APIKey.objects.get(key=key)
        if not api_key.is_valid():
            api_key.delete()
            return None
        return api_key
    except (ValueError, APIKey.DoesNotExist):
        return None


router = Router(auth=api_key_auth)


# Ping 测试的数据结构 ===
class TestConnectionIn(Schema):
    type: str
    provider: str = None
    api_key: str = None
    model: str = None


class EmbeddingIn(Schema):
    inputs: list[str] | None = None
    text: str = None
    provider: str = "siliconflow"
    model: str = "Qwen/Qwen3-Embedding-8B"
    api_key: str = None


# Ping 测试后端接口 ===
@router.post(
    "/test_connection", response={200: dict, 400: ErrorResponse, 408: ErrorResponse}
)
def test_connection(request, data: TestConnectionIn):
    start_time = time.time()

    try:
        if data.type == "provider":
            # 【修复点 1】：强制转为小写并去除空格，防止大小写不匹配
            provider = (data.provider or "").strip().lower()
            api_key = data.api_key
            headers = {"Authorization": f"Bearer {api_key}"}

            if provider == "deepseek":
                r = requests.get(
                    "https://api.deepseek.com/models", headers=headers, timeout=8
                )
            elif provider == "openai":
                r = requests.get(
                    "https://api.openai.com/v1/models", headers=headers, timeout=8
                )
            elif provider == "minimax":
                r = requests.get(
                    "https://api.minimax.chat/v1/models", headers=headers, timeout=8
                )
            elif provider == "siliconflow":
                r = requests.get(
                    "https://api.siliconflow.cn/v1/models", headers=headers, timeout=8
                )
            elif provider == "ollama":
                r = requests.get("http://localhost:11434/api/tags", timeout=4)
            else:
                return 400, {
                    "error": f"后端的 Ping 不支持此模型提供商: {data.provider}"
                }

            r.raise_for_status()

        elif data.type == "search":
            headers = {
                "Authorization": f"Bearer {data.api_key}",
                "Content-Type": "application/json",
            }
            payload = {
                "query": "test",
                "summary": True,
                "count": 1,
                "freshness": "noLimit",
            }
            r = requests.post(
                "https://api.bocha.cn/v1/web-search",
                headers=headers,
                json=payload,
                timeout=8,
            )
            r.raise_for_status()

        else:
            return 400, {"error": "未知的测试类型"}

        latency_ms = int((time.time() - start_time) * 1000)
        return 200, {"status": "ok", "latency_ms": latency_ms}

    except requests.exceptions.HTTPError as e:
        # 【修复点 2】：提取第三方机房返回的真实错误原因（比如余额不足、Key错误等）
        err_msg = str(e)
        try:
            err_json = e.response.json()
            if "error" in err_json:
                err_msg = str(err_json.get("error", err_json))
        except Exception:
            pass
        return 400, {
            "error": f"第三方拒绝访问 (HTTP {e.response.status_code}): {err_msg}"
        }

    except requests.exceptions.RequestException as e:
        return 408, {"error": f"网络不可达或超时，请检查服务器网络或代理设置"}


@router.get("/runtime-config", response={200: RuntimeConfigOut, 401: ErrorResponse})
def runtime_config(request):
    if not request.auth:
        return 401, {"error": "请先登录获取API Key"}

    provider_api_keys = {
        provider: services.resolve_provider_api_key(provider, None) or ""
        for provider in services.PROVIDER_API_KEY_ENV
    }

    return {
        "provider_api_keys": provider_api_keys,
        "web_search_api_key": services.resolve_web_search_api_key(None) or "",
    }


@router.post("/embeddings", response={200: dict, 400: ErrorResponse})
def create_embeddings(request, data: EmbeddingIn):
    provider = (data.provider or "siliconflow").strip().lower()
    if provider != "siliconflow":
        return 400, {"error": "当前仅支持 siliconflow 远程 embedding"}

    normalized_inputs: list[str] = []
    if isinstance(data.inputs, list):
        normalized_inputs.extend(
            str(item).strip() for item in data.inputs if str(item).strip()
        )

    single_text = (data.text or "").strip()
    if single_text:
        normalized_inputs.append(single_text)

    if not normalized_inputs:
        return 400, {"error": "请输入 text 或 inputs"}

    try:
        payload = services.create_remote_embeddings(
            inputs=normalized_inputs,
            model_name=data.model,
            provider_api_key=data.api_key,
        )
        return {"status": "ok", **payload}
    except ValueError as e:
        return 400, {"error": str(e)}
    except services.ProviderHttpError as e:
        message = e.detail.get("message") if isinstance(e.detail, dict) else None
        return 400, {"error": message or str(e)}


def _sse_line(payload: dict) -> str:
    return "data: " + json.dumps(payload, ensure_ascii=False) + "\n\n"


def _error_event(message: str, detail: dict | None = None) -> dict:
    payload = {
        "type": "error",
        "message": message,
        # 向后兼容旧前端字段
        "chunk": message,
    }
    if isinstance(detail, dict) and detail:
        payload["error_detail"] = detail
    return payload


def _sse_error_response(status_code: int, message: str) -> StreamingHttpResponse:
    return StreamingHttpResponse(
        _sse_line(_error_event(message)),
        status=status_code,
        content_type="text/event-stream",
    )


def _is_asgi_request(request) -> bool:
    # ASGIRequest 挂载了 scope，WSGIRequest 不会有该属性。
    return hasattr(request, "scope")


def _validate_office_archive(file_obj, filename: str) -> tuple[bool, str]:
    file_obj.seek(0)
    signature = file_obj.read(4)
    file_obj.seek(0)
    if signature != b"PK\x03\x04":
        return False, f"{filename} 文件结构不合法"

    with zipfile.ZipFile(file_obj) as archive:
        entries = archive.infolist()
        if len(entries) > settings.MAX_OFFICE_ARCHIVE_ENTRIES:
            return False, "Office 文件内条目过多，疑似压缩炸弹"

        total_uncompressed = sum(info.file_size for info in entries)
        if total_uncompressed > settings.MAX_OFFICE_UNCOMPRESSED_SIZE:
            return False, "Office 文件解压后体积过大，已拒绝"

    file_obj.seek(0)
    return True, ""


def clean_llm_reply(reply: str) -> str:
    """从模型原始回复中移除 <think>...</think> 标签块。"""
    return re.sub(r"<think>.*?</think>\s*", "", reply, flags=re.DOTALL).strip()


def _build_agent_cfg(base_llm: LlmConfig, agent_cfgs: dict, agent_id: str) -> LlmConfig:
    raw = agent_cfgs.get(agent_id) if isinstance(agent_cfgs, dict) else None
    if not isinstance(raw, dict):
        return base_llm
    return LlmConfig(
        provider=(raw.get("provider") or base_llm.provider or "ollama"),
        model=(raw.get("model") or base_llm.model),
        provider_api_key=(raw.get("provider_api_key") or base_llm.provider_api_key),
        embedding_mode=(raw.get("embedding_mode") or base_llm.embedding_mode),
        embedding_model=(raw.get("embedding_model") or base_llm.embedding_model),
    )


def _build_multi_agent_config(
    base_llm: LlmConfig, agent_cfgs: dict
) -> MultiAgentConfig:
    return MultiAgentConfig(
        rag=_build_agent_cfg(base_llm, agent_cfgs, "rag"),
        web=_build_agent_cfg(base_llm, agent_cfgs, "web"),
        synthesis=_build_agent_cfg(base_llm, agent_cfgs, "synthesis"),
    )


def _serialize_multi_agent_meta(agent_meta: dict | None) -> str:
    if not isinstance(agent_meta, dict) or not agent_meta:
        return ""

    try:
        payload = json.dumps(agent_meta, ensure_ascii=False, separators=(",", ":"))
    except TypeError:
        return ""

    return f"【MULTI_AGENT_META】{payload}【/MULTI_AGENT_META】"


def _render_context(
    history_for_llm,
    user_input: str,
    final_reply: str,
    current_agent_meta: dict | None = None,
) -> str:
    lines = []
    for msg in history_for_llm:
        if msg.get("role") == "user":
            lines.append(f"用户：{msg.get('content', '')}")
        elif msg.get("role") == "assistant":
            lines.append(f"回复：{msg.get('content', '')}")
            agent_meta = msg.get("agent_meta")
            meta_line = _serialize_multi_agent_meta(agent_meta)
            if meta_line:
                lines.append(meta_line)

    lines.append(f"用户：{user_input}")
    lines.append(f"回复：{final_reply}")
    meta_line = _serialize_multi_agent_meta(current_agent_meta)
    if meta_line:
        lines.append(meta_line)
    return "\n".join(lines).strip()


def _build_query_filters(
    db_type: str,
    risk_level: str,
    source: str,
    cve_id: str,
    ioc_value: str,
    tag: str,
    mitre_attack_id: str,
) -> dict:
    return {
        "db_type": db_type,
        "risk_level": risk_level,
        "source": source,
        "cve_id": cve_id,
        "ioc_value": ioc_value,
        "tag": tag,
        "mitre_attack_id": mitre_attack_id,
    }


def _split_csv(value: str) -> list[str]:
    text = (value or "").strip()
    if not text:
        return []
    return [item.strip() for item in text.replace(";", ",").split(",") if item.strip()]


@api.post("/login", response={200: LoginOut, 400: ErrorResponse, 403: ErrorResponse})
def login(request, data: LoginIn):
    username = data.username.strip()
    password = data.password.strip()
    if not username or not password:
        return 400, {"error": "用户名和密码不能为空"}
    if password != settings.AUTH_PASSWORD:
        return 403, {"error": "密码错误"}
    key = services.create_api_key(username)
    return {"api_key": key, "expiry": settings.TOKEN_EXPIRY_SECONDS}


@api.get("/health")
def health(request):
    return {
        "status": "ok",
        "service": "deepsoc-api",
    }


@api.get("/ready")
def ready(request):
    connection.ensure_connection()
    log_ready = services.log_system is not None
    status_code = 200 if log_ready else 503
    return status_code, {
        "status": "ready" if log_ready else "degraded",
        "db": "ok",
        "vector": "ok" if log_ready else "unavailable",
    }


@router.post("/chat")
def chat(request, data: ChatIn):
    if not request.auth:
        return _sse_error_response(401, "请先登录获取API Key")

    session_id = data.session_id.strip() or "默认对话"
    user_input = data.user_input.strip()
    if not user_input:
        return _sse_error_response(400, "请输入消息内容")

    user = request.auth
    session = get_or_create_session(session_id, user)

    use_db_search = data.use_db_search
    use_web_search = data.use_web_search
    selected_model = data.model_name
    selected_provider = (data.provider or "ollama").strip().lower()
    selected_embedding_mode = (data.embedding_mode or "local").strip().lower()
    selected_embedding_model = data.embedding_model
    provider_api_key = data.provider_api_key
    web_search_api_key = data.web_search_api_key
    mode = (data.mode or "").strip().lower() or "single"

    logger.info(
        "搜索选项 - 数据库: %s, 联网: %s, provider: %s, model: %s, embedding_mode: %s, embedding_model: %s",
        use_db_search,
        use_web_search,
        selected_provider,
        selected_model or "default",
        selected_embedding_mode,
        selected_embedding_model or "default",
    )

    if data.context is not None:
        history_for_llm = list(data.context)
        is_regeneration = False
        logger.info(
            "使用前端提供的 context (会话: %s, 条数: %s)",
            session_id,
            len(history_for_llm),
        )
    else:
        conversation_history = session.get_conversation_history()
        history_for_llm = conversation_history
        is_regeneration = False

        if (
            len(conversation_history) >= 2
            and conversation_history[-1].get("role") == "assistant"
            and conversation_history[-2].get("role") == "user"
            and conversation_history[-2].get("content") == user_input
        ):
            logger.info("检测到重新生成 (会话: %s)", session_id)
            history_for_llm = conversation_history[:-2]
            is_regeneration = True
        elif (
            len(conversation_history) >= 1
            and conversation_history[-1].get("role") == "user"
            and conversation_history[-1].get("content") == user_input
        ):
            logger.info("检测到对失败消息的重新生成 (会话: %s)", session_id)
            history_for_llm = conversation_history[:-1]
            is_regeneration = True

    def stream_generator() -> Generator[str, None, None]:
        full_clean_reply = ""
        agent_state = {
            "rag": {"status": "idle", "content": "", "error": "", "errorDetail": None},
            "web": {"status": "idle", "content": "", "error": "", "errorDetail": None},
        }

        # 先发送一个轻量心跳，确保客户端尽快进入流式读取状态。
        yield _sse_line({"type": "ping"})

        def _update_agent_state_from_event(evt: dict) -> None:
            agent_id = evt.get("agent_id")
            if agent_id not in agent_state:
                return

            if evt.get("type") == "agent_chunk":
                agent_state[agent_id]["content"] += (
                    evt.get("content", evt.get("chunk", "")) or ""
                )
                return

            if evt.get("type") == "agent_status":
                agent_state[agent_id]["status"] = (
                    evt.get("status") or agent_state[agent_id]["status"]
                )
                if evt.get("status") == "error":
                    agent_state[agent_id]["error"] = (
                        evt.get("error") or evt.get("message") or ""
                    )
                    agent_state[agent_id]["errorDetail"] = evt.get("error_detail")
                elif evt.get("status") == "done" and not agent_state[agent_id]["error"]:
                    agent_state[agent_id]["errorDetail"] = None

        try:
            if mode != "multi_agent":
                for raw_chunk in model_api_call(
                    user_input,
                    history_for_llm,
                    use_db_search,
                    use_web_search,
                    model_name=selected_model,
                    provider=selected_provider,
                    embedding_mode=selected_embedding_mode,
                    embedding_model=selected_embedding_model,
                    provider_api_key=provider_api_key,
                    web_search_api_key=web_search_api_key,
                ):
                    if not raw_chunk:
                        continue

                    if isinstance(raw_chunk, dict):
                        if raw_chunk.get("type") == "error":
                            msg = (
                                raw_chunk.get("message")
                                or raw_chunk.get("chunk")
                                or "模型调用失败"
                            ).strip()
                            detail = raw_chunk.get("error_detail")
                            yield _sse_line(
                                _error_event(
                                    msg, detail if isinstance(detail, dict) else None
                                )
                            )
                        else:
                            yield _sse_line(raw_chunk)
                        continue

                    yield _sse_line({"type": "content", "chunk": raw_chunk})
                    full_clean_reply += raw_chunk
            else:
                base_llm = LlmConfig(
                    provider=selected_provider,
                    model=selected_model,
                    provider_api_key=provider_api_key,
                    embedding_mode=selected_embedding_mode,
                    embedding_model=selected_embedding_model,
                )
                cfg = _build_multi_agent_config(base_llm, data.agent_configs or {})
                orch = Orchestrator()

                for evt in orch.run_stream(
                    OrchestratorInput(
                        query=user_input,
                        history=history_for_llm,
                        enable_rag=use_db_search,
                        enable_web=use_web_search,
                        web_search_api_key=web_search_api_key,
                    ),
                    cfg,
                ):
                    _update_agent_state_from_event(evt)
                    yield _sse_line(evt)
                    if (
                        evt.get("type") == "agent_chunk"
                        and evt.get("agent_id") == "synthesis"
                    ):
                        full_clean_reply += evt.get("content", "")

            final_save = full_clean_reply.strip()
            if not final_save:
                logger.warning("会话 %s 未收到有效模型输出，跳过写入", session_id)
                return

            current_agent_meta = None
            if mode == "multi_agent":
                current_agent_meta = agent_state

            if data.context is not None or is_regeneration:
                session.context = _render_context(
                    history_for_llm, user_input, final_save, current_agent_meta
                )
                session.save()
            else:
                if current_agent_meta:
                    session.update_context(
                        user_input,
                        f"{final_save}\n{_serialize_multi_agent_meta(current_agent_meta)}",
                    )
                else:
                    session.update_context(user_input, final_save)

            logger.info("会话 %s 已更新 (用户: %s)", session_id, user.user)
        except Exception as e:
            logger.error("流式生成失败: %s", e)
            if mode == "multi_agent":
                detail = getattr(e, "detail", None)
                payload = {
                    "type": "agent_status",
                    "agent_id": "synthesis",
                    "status": "error",
                    "error": f"流处理失败: {e}",
                }
                if isinstance(detail, dict):
                    payload["error_detail"] = detail
                yield _sse_line(payload)
            else:
                yield _sse_line(_error_event(f"流处理失败: {e}"))
        finally:
            yield _sse_line({"type": "done"})

    if _is_asgi_request(request):

        async def async_stream_generator():
            iterator = iter(stream_generator())
            while True:
                chunk = await sync_to_async(next, thread_sensitive=True)(iterator, None)
                if chunk is None:
                    break
                yield chunk

        streaming_content = async_stream_generator()
    else:
        streaming_content = stream_generator()

    response = StreamingHttpResponse(streaming_content, content_type="text/event-stream")
    response["X-Accel-Buffering"] = "no"
    response["Cache-Control"] = "no-cache, no-transform"
    return response


@router.get("/sessions", response={200: SessionListOut})
def sessions(request):
    return {"sessions": list_user_sessions(request.auth)}


@router.get("/history", response={200: HistoryOut})
def history(request, session_id: str = "默认对话"):
    processed_session_id = session_id.strip() or "默认对话"
    session = services.get_or_create_session(processed_session_id, request.auth)
    return {"history": session.context}


@router.delete("/history", response={200: dict})
def clear_history(request, session_id: str = "默认对话"):
    processed_session_id = session_id.strip() or "默认对话"
    session = services.get_or_create_session(processed_session_id, request.auth)
    session.clear_context()
    return {"message": "历史记录已清空"}


@router.delete("/session", response={200: dict})
def delete_session(request, session_id: str = "默认对话"):
    normalized_session_id, deleted = delete_conversation_session(request.auth, session_id)
    return {
        "message": "ok",
        "session_id": normalized_session_id,
        "deleted": deleted,
    }


@router.post("/session/rename", response={200: dict, 400: ErrorResponse})
def rename_session(request, data: SessionRenameIn):
    """重命名会话：更新数据库中的 session_id，与前端会话列表一致。"""
    try:
        old_id, new_id = rename_conversation_session(
            request.auth, data.old_session_id, data.new_session_id
        )
        return {
            "message": "ok",
            "old_session_id": old_id,
            "new_session_id": new_id,
        }
    except ValueError as e:
        return 400, {"error": str(e)}


@router.get("/dashboard/stats", response={200: dict})
def get_dashboard_stats(request):
    """聚合 data/log 下真实 CSV 日志，返回大屏图表与拓扑数据。"""
    return build_dashboard_stats()


@router.get("/query/logs", response={200: dict})
def query_logs(
    request,
    q: str = "",
    page: int = 1,
    page_size: int = 20,
    db_type: str = "",
    risk_level: str = "",
    source: str = "",
    cve_id: str = "",
    ioc_value: str = "",
    tag: str = "",
    mitre_attack_id: str = "",
    start_time: str = "",
    end_time: str = "",
    sort_by: str = "fetched_at",
    sort_order: str = "desc",
):
    filters = _build_query_filters(
        db_type=db_type,
        risk_level=risk_level,
        source=source,
        cve_id=cve_id,
        ioc_value=ioc_value,
        tag=tag,
        mitre_attack_id=mitre_attack_id,
    )
    return list_query_records(
        query=q,
        page=page,
        page_size=page_size,
        filters=filters,
        start_time=start_time,
        end_time=end_time,
        sort_by=sort_by,
        sort_order=sort_order,
    )


@router.get("/query/logs/{record_id}", response={200: dict, 404: ErrorResponse})
def query_log_detail(request, record_id: str):
    item = get_query_record_detail(record_id)
    if not item:
        return 404, {"error": "记录不存在"}
    return item


@router.get("/query/facets", response={200: dict})
def query_facets(
    request,
    q: str = "",
    db_type: str = "",
    risk_level: str = "",
    source: str = "",
    cve_id: str = "",
    ioc_value: str = "",
    tag: str = "",
    mitre_attack_id: str = "",
    start_time: str = "",
    end_time: str = "",
):
    filters = _build_query_filters(
        db_type=db_type,
        risk_level=risk_level,
        source=source,
        cve_id=cve_id,
        ioc_value=ioc_value,
        tag=tag,
        mitre_attack_id=mitre_attack_id,
    )
    return get_query_facets(
        query=q,
        filters=filters,
        start_time=start_time,
        end_time=end_time,
    )


@router.get("/query/export")
def query_export(
    request,
    export_format: str = "csv",
    q: str = "",
    db_type: str = "",
    risk_level: str = "",
    source: str = "",
    cve_id: str = "",
    ioc_value: str = "",
    tag: str = "",
    mitre_attack_id: str = "",
    start_time: str = "",
    end_time: str = "",
    sort_by: str = "fetched_at",
    sort_order: str = "desc",
    export_scope: str = "all",
    page: int = 1,
    page_size: int = 20,
    fields: str = "",
    include_details: bool = False,
    filename_prefix: str = "deepsoc_query",
):
    filters = _build_query_filters(
        db_type=db_type,
        risk_level=risk_level,
        source=source,
        cve_id=cve_id,
        ioc_value=ioc_value,
        tag=tag,
        mitre_attack_id=mitre_attack_id,
    )
    selected_fields = _split_csv(fields)
    content, content_type, filename = export_query_records(
        export_format=export_format,
        query=q,
        filters=filters,
        start_time=start_time,
        end_time=end_time,
        sort_by=sort_by,
        sort_order=sort_order,
        export_scope=export_scope,
        page=page,
        page_size=page_size,
        selected_fields=selected_fields,
        include_details=include_details,
        filename_prefix=filename_prefix,
    )

    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@router.post("/upload_file")
def upload_file(request, file: NinjaUploadedFile = File(...)):
    if not request.auth:
        return 401, {"error": "请先登录获取API Key"}

    raw_name = (file.name or "").strip()
    safe_name = os.path.basename(raw_name)
    if not safe_name:
        return 400, {"error": "文件名不能为空"}

    filename = safe_name.lower()
    _, ext = os.path.splitext(filename)
    allowed_exts = {".txt", ".docx", ".xlsx"}
    if ext not in allowed_exts:
        return 400, {"error": "不支持的文件类型，仅支持 .txt / .docx / .xlsx"}

    file_size = int(getattr(file, "size", 0) or 0)
    if file_size <= 0:
        return 400, {"error": "上传文件为空"}
    if file_size > settings.UPLOAD_MAX_BYTES:
        max_mb = settings.UPLOAD_MAX_BYTES // (1024 * 1024)
        return 400, {"error": f"文件过大，最大允许 {max_mb}MB"}

    try:
        if ext == ".txt":
            content_bytes = file.read()
            try:
                text = content_bytes.decode("utf-8")
            except Exception:
                text = content_bytes.decode("gbk", errors="ignore")
            return {"text": text}

        if ext == ".docx":
            try:
                from docx import Document
            except Exception:
                return 400, {"error": "缺少依赖：请安装 python-docx"}

            is_valid, reason = _validate_office_archive(file.file, safe_name)
            if not is_valid:
                return 400, {"error": reason}

            document = Document(file)
            paragraphs = [p.text for p in document.paragraphs if p.text]
            text = "\n".join(paragraphs)
            return {"text": text}

        if ext == ".xlsx":
            try:
                import openpyxl
            except Exception:
                return 400, {"error": "缺少依赖：请安装 openpyxl"}

            is_valid, reason = _validate_office_archive(file.file, safe_name)
            if not is_valid:
                return 400, {"error": reason}

            wb = openpyxl.load_workbook(file, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"# 工作表: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v) for v in row]
                    lines.append("\t".join(cells))
            text = "\n".join(lines)
            return {"text": text}

        return 400, {"error": "不支持的文件类型，仅支持 .txt / .docx / .xlsx"}

    except Exception as e:
        logger.error("文件解析失败: %s", e)
        return 400, {"error": f"文件解析失败: {e}"}


api.add_router("", router)
